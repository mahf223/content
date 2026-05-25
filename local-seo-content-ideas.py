"""Generate an Excel sheet with trending Local SEO content topic ideas for 2026."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Local SEO Content Ideas"

# ---- Styling ----
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
category_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
category_font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
cell_font = Font(name="Calibri", size=11)
top_pick_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

thin = Side(border_style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---- Headers ----
headers = [
    "#",
    "Category",
    "Blog Topic Idea",
    "Primary Keyword",
    "Search Intent",
    "Difficulty",
    "Trending Score (1-10)",
    "Suggested Word Count",
    "Top Pick",
]
ws.append(headers)

for col_idx, _ in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

ws.row_dimensions[1].height = 32

# ---- Data ----
# (category, topic, keyword, intent, difficulty, trending, words, top_pick)
data = [
    # AI & Search Trends
    ("AI & Search Trends", "How Google AI Overviews Are Changing Local Search Results in 2026", "google ai overviews local seo", "Informational", "Medium", 10, 1800, "Yes"),
    ("AI & Search Trends", "Local SEO for ChatGPT, Perplexity & Gemini: Get Your Business Cited by AI", "local seo for chatgpt", "Informational", "Medium", 10, 2000, "Yes"),
    ("AI & Search Trends", "GEO (Generative Engine Optimization) vs Local SEO: What's the Difference?", "geo vs local seo", "Informational", "Low", 10, 1500, "Yes"),
    ("AI & Search Trends", "How to Optimize Your GBP for AI-Powered 'Near Me' Searches", "optimize gbp for ai search", "How-to", "Medium", 9, 1800, "No"),
    ("AI & Search Trends", "Voice Search Local SEO: Ranking for 'Hey Google, Find a [Service] Near Me'", "voice search local seo", "How-to", "Medium", 8, 1600, "No"),

    # Google Business Profile
    ("Google Business Profile", "15 GBP Features Most Local Businesses Aren't Using in 2026", "google business profile features", "Informational", "Low", 9, 2200, "No"),
    ("Google Business Profile", "Google Business Profile Posts: A Weekly Strategy That Drives Calls", "google business profile posts strategy", "How-to", "Low", 8, 1500, "No"),
    ("Google Business Profile", "How to Recover from a Google Business Profile Suspension", "gbp suspension recovery", "How-to", "Medium", 8, 1800, "No"),
    ("Google Business Profile", "GBP Q&A Section: How to Use It as a Free SEO Asset", "google business profile q&a seo", "How-to", "Low", 7, 1200, "No"),
    ("Google Business Profile", "New Google Business Profile Updates 2026: What Marketers Need to Know", "google business profile updates 2026", "Informational", "Low", 10, 1400, "Yes"),

    # Reviews & Reputation
    ("Reviews & Reputation", "How to Get 100 Google Reviews in 90 Days (Without Violating Guidelines)", "how to get more google reviews", "How-to", "Medium", 9, 1800, "No"),
    ("Reviews & Reputation", "Responding to Negative Reviews: 10 Templates That Win Customers Back", "respond to negative google reviews", "How-to", "Low", 8, 1600, "No"),
    ("Reviews & Reputation", "Why Review Velocity Matters More Than Total Review Count", "google review velocity seo", "Informational", "Low", 7, 1200, "No"),
    ("Reviews & Reputation", "AI-Generated Fake Reviews: How to Spot and Report Them", "ai fake google reviews", "Informational", "Low", 9, 1500, "No"),

    # Map Pack & Rankings
    ("Map Pack & Rankings", "How to Rank in the Google Local 3-Pack: A 2026 Checklist", "rank in google local 3-pack", "How-to", "High", 9, 2500, "No"),
    ("Map Pack & Rankings", "Why Your Business Disappears from the Map Pack on Mobile (Fix It)", "business disappeared from google maps", "How-to", "Medium", 8, 1500, "No"),
    ("Map Pack & Rankings", "Local SEO Ranking Factors That Actually Moved the Needle in 2026", "local seo ranking factors 2026", "Informational", "Medium", 9, 2200, "No"),
    ("Map Pack & Rankings", "Proximity vs Relevance: Which Local Ranking Factor Matters More Now?", "proximity vs relevance local seo", "Informational", "Low", 7, 1300, "No"),

    # Content & On-Page
    ("Content & On-Page", "Hyperlocal Blogging: How to Write City + Neighborhood Pages That Rank", "hyperlocal seo content", "How-to", "Medium", 8, 2000, "No"),
    ("Content & On-Page", "Local Landing Page Template: 12 Sections Every Service Page Needs", "local landing page template", "How-to", "Medium", 8, 1800, "No"),
    ("Content & On-Page", "How to Use Schema Markup for Local Businesses (with Code Examples)", "local business schema markup", "How-to", "Medium", 8, 2000, "No"),
    ("Content & On-Page", "'Near Me' Keywords: How to Find and Target Them in 2026", "near me keywords local seo", "How-to", "Low", 8, 1500, "No"),
    ("Content & On-Page", "Seasonal Local SEO: Content Calendar for Service Businesses", "seasonal local seo content", "How-to", "Low", 7, 1600, "No"),

    # Citations & Links
    ("Citations & Links", "NAP Consistency in 2026: Does It Still Matter for Local Rankings?", "nap consistency local seo 2026", "Informational", "Low", 7, 1400, "No"),
    ("Citations & Links", "15 Free High-DA Local Citation Sites You Probably Missed", "free local citation sites", "Listicle", "Low", 8, 1800, "No"),
    ("Citations & Links", "How to Build Local Backlinks Without Cold Outreach", "local backlinks without outreach", "How-to", "Medium", 8, 2000, "No"),
    ("Citations & Links", "Sponsoring Local Events for SEO: Real ROI or Waste of Money?", "local event sponsorship seo", "Informational", "Low", 6, 1400, "No"),

    # Niche / Industry-Specific
    ("Niche / Industry", "Local SEO for Plumbers: 7 Strategies That Generate Daily Leads", "local seo for plumbers", "How-to", "Medium", 9, 2200, "Yes"),
    ("Niche / Industry", "Local SEO for Dentists in 2026: Beating Big Chains in Your City", "local seo for dentists", "How-to", "High", 9, 2500, "No"),
    ("Niche / Industry", "Real Estate Agent Local SEO: Ranking for '[City] Homes for Sale'", "real estate agent local seo", "How-to", "High", 8, 2200, "No"),
    ("Niche / Industry", "Restaurant Local SEO: Show Up in Google Maps Food Searches", "restaurant local seo", "How-to", "Medium", 9, 2000, "No"),
    ("Niche / Industry", "Lawyer Local SEO: Compliance + Ranking Strategies", "local seo for lawyers", "How-to", "High", 8, 2500, "No"),

    # Tools & Tutorials
    ("Tools & Tutorials", "Free Local SEO Audit: 25-Point Checklist with Screenshots", "free local seo audit checklist", "How-to", "Medium", 9, 2400, "No"),
    ("Tools & Tutorials", "Best Local Rank Tracking Tools Compared (2026 Edition)", "best local rank tracking tools 2026", "Listicle", "Medium", 8, 2000, "No"),
    ("Tools & Tutorials", "How to Set Up Local SEO Tracking in GA4 + Google Search Console", "local seo tracking ga4", "How-to", "Medium", 7, 1800, "No"),
]

start_row = 2
for idx, row in enumerate(data, start=1):
    category, topic, keyword, intent, difficulty, trending, words, top_pick = row
    excel_row = [idx, category, topic, keyword, intent, difficulty, trending, words, top_pick]
    ws.append(excel_row)

    current_row = start_row + idx - 1
    is_top = top_pick == "Yes"

    for col_idx in range(1, len(excel_row) + 1):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.font = cell_font
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if col_idx in (1, 5, 6, 7, 8, 9):
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if is_top:
            cell.fill = top_pick_fill

# ---- Column widths ----
widths = {
    "A": 5,    # #
    "B": 24,   # Category
    "C": 70,   # Topic
    "D": 32,   # Primary Keyword
    "E": 16,   # Intent
    "F": 12,   # Difficulty
    "G": 14,   # Trending
    "H": 14,   # Word Count
    "I": 10,   # Top Pick
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

# Freeze header row
ws.freeze_panes = "A2"

# ---- Second sheet: Top 5 Picks summary ----
ws2 = wb.create_sheet("Top 5 Picks")
ws2.append(["Rank", "Topic", "Why It's Trending"])
top5 = [
    (1, "Local SEO for ChatGPT, Perplexity & Gemini", "AI search engines are reshaping how people find local businesses; very low competition right now."),
    (2, "How Google AI Overviews Are Changing Local Search", "AI Overviews are eating clicks from the map pack; high search volume in 2026."),
    (3, "GEO vs Local SEO: What's the Difference?", "New term marketers are searching for; ranks fast due to low competition."),
    (4, "New Google Business Profile Updates 2026", "Always trends when Google releases GBP changes; strong recurring traffic."),
    (5, "Local SEO for Plumbers / Dentists / [Your Niche]", "Niche local SEO posts attract agency clients and high-intent leads."),
]
for row in top5:
    ws2.append(row)

for col_idx in range(1, 4):
    cell = ws2.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

for r in range(2, 7):
    for c in range(1, 4):
        cell = ws2.cell(row=r, column=c)
        cell.font = cell_font
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if c == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")

ws2.column_dimensions["A"].width = 8
ws2.column_dimensions["B"].width = 55
ws2.column_dimensions["C"].width = 70
ws2.row_dimensions[1].height = 28
ws2.freeze_panes = "A2"

# ---- Save ----
output = "/projects/sandbox/content/Local-SEO-Content-Ideas-2026.xlsx"
wb.save(output)
print(f"Saved: {output}")
print(f"Total rows: {len(data)}")
